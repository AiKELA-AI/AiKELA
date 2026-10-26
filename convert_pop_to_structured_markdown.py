#!/usr/bin/env python3
"""Convert Programme on a Plate files to structure-preserving Markdown.

Legacy Word files use macOS textutil's HTML export so table rows and cells can
be retained. The downloaded originals are kept separately under src/ while
Markdown output defaults to the POP folder.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import subprocess
import sys
import tempfile
from collections import Counter
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parent
POP_ROOT = ROOT / "docs" / "pop"
SOURCE = POP_ROOT / "src"
DESTINATION = POP_ROOT
SOURCE_PAGE = "https://scoutwiki.scouts.org.za/wiki/Programmes_on_a_Plate"
SUPPORTED = {".doc", ".docx", ".pdf", ".xlsx"}


def clean(text: str) -> str:
    return re.sub(r"\s+", " ", text.replace("\xa0", " ")).strip()


def md_cell(text: str) -> str:
    return text.replace("|", r"\|").replace("\n", "<br>")


class WordHtmlParser(HTMLParser):
    """Read textutil HTML in source order, preserving table rows and cells."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.blocks: list[tuple[str, Any]] = []
        self.table: list[list[dict[str, Any]]] | None = None
        self.row: list[dict[str, Any]] | None = None
        self.cell: dict[str, Any] | None = None
        self.block_tag: str | None = None
        self.block_text: list[str] = []
        self.metadata: dict[str, str] = {}
        self.list_depth = 0
        self.ignored_tags: list[str] = []

    def flush_block(self) -> None:
        value = clean("".join(self.block_text))
        if value:
            self.blocks.append((self.block_tag or "p", value))
        self.block_tag = None
        self.block_text = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        values = dict(attrs)
        if tag == "meta" and values.get("name") and values.get("content"):
            self.metadata[values["name"].lower()] = clean(values["content"] or "")
        if tag in {"head", "style", "script", "title"}:
            self.ignored_tags.append(tag)
            return
        if self.ignored_tags:
            return
        if tag == "table":
            self.flush_block()
            self.table = []
            self.blocks.append(("table", self.table))
        elif tag == "tr" and self.table is not None:
            self.row = []
            self.table.append(self.row)
        elif tag in {"td", "th"} and self.row is not None:
            self.cell = {"parts": [], "colspan": int(values.get("colspan") or 1)}
            self.row.append(self.cell)
        elif tag in {"p", "h1", "h2", "h3", "h4", "h5", "h6", "li"} and self.table is None:
            self.flush_block()
            self.block_tag = tag
        elif tag == "br":
            if self.cell is not None:
                self.cell["parts"].append(" ")
            else:
                self.block_text.append(" ")
        elif tag in {"ul", "ol"} and self.table is None:
            self.list_depth += 1

    def handle_endtag(self, tag: str) -> None:
        if tag in self.ignored_tags:
            self.ignored_tags.remove(tag)
            return
        if self.ignored_tags:
            return
        if tag in {"td", "th"}:
            self.cell = None
        elif tag == "tr":
            self.row = None
        elif tag == "table":
            self.table = None
        elif tag in {"p", "h1", "h2", "h3", "h4", "h5", "h6", "li"}:
            if self.cell is None:
                self.flush_block()
        elif tag in {"ul", "ol"}:
            self.list_depth = max(0, self.list_depth - 1)

    def handle_data(self, data: str) -> None:
        if self.ignored_tags:
            return
        if self.cell is not None:
            self.cell["parts"].append(data)
        else:
            self.block_text.append(data)


def convert_word(path: Path) -> tuple[str, dict[str, str]]:
    result = subprocess.run(
        ["textutil", "-convert", "html", "-stdout", str(path)],
        check=False,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    if result.returncode:
        detail = result.stderr.strip() or f"textutil exited with status {result.returncode}"
        raise RuntimeError(detail)

    parser = WordHtmlParser()
    parser.feed(result.stdout.replace("\x00", ""))
    parser.flush_block()
    parts: list[str] = []
    for kind, value in parser.blocks:
        if kind == "table":
            rendered = render_table(value)
            if rendered:
                parts.append(rendered)
        elif kind == "li":
            parts.append(f"- {value}")
        elif kind.startswith("h") and len(kind) == 2 and kind[1].isdigit():
            parts.append("#" * min(int(kind[1]) + 1, 6) + f" {value}")
        else:
            parts.append(value)
    return "\n\n".join(parts), parser.metadata


def render_table(raw_rows: list[list[dict[str, Any]]]) -> str:
    rows: list[list[str]] = []
    spans: list[list[int]] = []
    for row in raw_rows:
        values = [clean("".join(cell["parts"])) for cell in row]
        if any(values):
            rows.append(values)
            spans.append([cell["colspan"] for cell in row])
    if not rows:
        return ""

    width = max(len(row) for row in rows)
    # A single full-width row in an activity plan is a contextual programme note.
    headers = [value.casefold().strip(" :") for value in rows[0]]
    is_activity_table = {"activity", "equipment"}.issubset(set(headers))
    if is_activity_table:
        header = rows[0] + [""] * (width - len(rows[0]))
        data_rows = []
        data_spans = []
        normalized_header = [value.strip(" :/").casefold() for value in header]
        for row, row_span in zip(rows[1:], spans[1:]):
            normalized_row = [value.strip(" :/").casefold() for value in row]
            if len(row) == width and normalized_row == normalized_header:
                continue  # Word often repeats the column headings after a page break.
            data_rows.append(row)
            data_spans.append(row_span)
    else:
        header = [f"Column {index + 1}" for index in range(width)]
        data_rows = rows
        data_spans = spans

    output = [
        "| " + " | ".join(md_cell(item) for item in header) + " |",
        "| " + " | ".join("---" for _ in range(width)) + " |",
    ]
    rendered_rows: list[list[str]] = []
    for row, row_spans in zip(data_rows, data_spans):
        if is_activity_table and len(row) == 1:
            note = row[0]
            row = [""] * width
            note_column = min(2, width - 1)
            row[note_column] = f"**Programme note:** {note}"
        else:
            row = row + [""] * (width - len(row))
        if (
            is_activity_table
            and len(row) > 2
            and not row[0]
            and not row[1]
            and row[2]
            and row[2][0].islower()
            and rendered_rows
            and len(rendered_rows[-1]) > 2
            and rendered_rows[-1][1]
            and not any(row[3:])
        ):
            rendered_rows[-1][2] = (rendered_rows[-1][2] + " " + row[2]).strip()
            continue  # Continue a wrapped activity row across a page boundary.
        rendered_rows.append(row)
    for row in rendered_rows:
        output.append("| " + " | ".join(md_cell(item) for item in row) + " |")
    return "\n".join(output)


def convert_pdf(path: Path) -> tuple[str, dict[str, str]]:
    reader = PdfReader(str(path))
    parts = []
    for number, page in enumerate(reader.pages, 1):
        # Layout mode keeps spacing and visual columns more faithfully than the
        # default reading-order extractor, which is useful for programme PDFs.
        text = (page.extract_text(extraction_mode="layout") or "").strip()
        if text:
            parts.append(f"## Page {number}\n\n{text}")
    if not parts:
        raise RuntimeError("PDF has no extractable text; refusing to write misleading empty Markdown")
    return "\n\n".join(parts), {"pages": str(len(reader.pages))}


def convert_xlsx(path: Path) -> tuple[str, dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    parts = []
    try:
        for sheet in workbook.worksheets:
            rows = [
                ["" if value is None else clean(str(value)) for value in row]
                for row in sheet.iter_rows(values_only=True)
            ]
            while rows and not any(rows[-1]):
                rows.pop()
            if not rows:
                continue
            width = max(len(row) for row in rows)
            rows = [row + [""] * (width - len(row)) for row in rows]
            parts.extend([f"## Sheet: {sheet.title}", "", render_plain_table(rows)])
    finally:
        workbook.close()
    if not parts:
        raise RuntimeError("Workbook has no non-empty sheets")
    return "\n".join(parts), {}


def render_plain_table(rows: list[list[str]]) -> str:
    width = max(len(row) for row in rows)
    padded = [row + [""] * (width - len(row)) for row in rows]
    # Spreadsheet first rows often contain data rather than a true heading row.
    output = [
        "| " + " | ".join(f"Column {i + 1}" for i in range(width)) + " |",
        "| " + " | ".join("---" for _ in range(width)) + " |",
    ]
    output.extend("| " + " | ".join(md_cell(value) for value in row) + " |" for row in padded)
    return "\n".join(output)


def convert(path: Path) -> tuple[str, dict[str, str]]:
    suffix = path.suffix.lower()
    if suffix in {".doc", ".docx"}:
        return convert_word(path)
    if suffix == ".pdf":
        return convert_pdf(path)
    if suffix == ".xlsx":
        return convert_xlsx(path)
    raise ValueError(f"Unsupported file type: {suffix}")


def build_markdown(path: Path, body: str, details: dict[str, str]) -> str:
    title = re.sub(r"[_-]+", " ", path.stem).strip()
    lines = [
        f"# {title}",
        "",
        f"- Original file: [`{path.name}`]({SOURCE_PAGE})",
        f"- Source: [Programme on a Plate, Scouts South Africa Wiki]({SOURCE_PAGE})",
    ]
    # textutil's legacy Word modification timestamp can differ from the file's OLE metadata.
    for source_key, label in (("author", "Author"), ("creationtime", "Created")):
        value = details.get(source_key)
        if value:
            lines.append(f"- {label}: {value}")
    if details.get("pages"):
        lines.append(f"- Pages: {details['pages']}")
    lines.extend(["", "---", "", body.strip(), ""])
    return "\n".join(lines)


def main() -> int:
    cli = argparse.ArgumentParser(description=__doc__)
    cli.add_argument("--source", type=Path, default=SOURCE)
    cli.add_argument("--destination", type=Path, default=DESTINATION)
    cli.add_argument("--only", help="Convert one source filename (or stem) for a focused check")
    args = cli.parse_args()
    source = args.source.expanduser().resolve()
    destination = args.destination.expanduser().resolve()

    if not source.is_dir():
        print(f"Source folder not found: {source}", file=sys.stderr)
        return 2
    files = sorted(p for p in source.iterdir() if p.is_file() and p.suffix.lower() in SUPPORTED)
    if args.only:
        files = [p for p in files if p.name == args.only or p.stem == args.only]
        if not files:
            print(f"No supported source file matched: {args.only}", file=sys.stderr)
            return 2

    destination.mkdir(parents=True, exist_ok=True)
    errors: list[dict[str, str]] = []
    converted = 0
    for path in files:
        target = destination / f"{path.stem}.md"
        try:
            body, details = convert(path)
            if not body.strip():
                raise RuntimeError("Conversion returned no content; refusing to write empty Markdown")
            document = build_markdown(path, body, details)
            # Replace only this converter's output, and only after conversion succeeded.
            with tempfile.NamedTemporaryFile(
                "w", encoding="utf-8", dir=destination, prefix=f".{path.stem}.", suffix=".tmp", delete=False
            ) as handle:
                temporary = Path(handle.name)
                handle.write(document)
            os.replace(temporary, target)
            converted += 1
            print(f"OK    {path.name} -> {target.name}")
        except Exception as error:
            errors.append({"source": path.name, "error": f"{type(error).__name__}: {error}"})
            print(f"FAIL  {path.name}: {errors[-1]['error']}")

    log = destination / "conversion-errors.json"
    log.write_text(json.dumps(errors, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    counts = Counter(path.suffix.lower() for path in files)
    print(f"\nFinished: {converted}/{len(files)} converted, {len(errors)} failed.")
    print("Source types: " + ", ".join(f"{ext}={count}" for ext, count in sorted(counts.items())))
    print(f"Markdown output: {destination}")
    print(f"Failure log: {log}")
    return 1 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
