from pathlib import Path
from subprocess import run, PIPE
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
import xlrd
import re

SOURCE = Path.home() / "cub-ai" / "docs" / "programmes-on-a-plate"
DEST = SOURCE / "markdown"

SUPPORTED = {
    ".pdf", ".doc", ".docx", ".xls", ".xlsx",
    ".ppt", ".pptx", ".txt", ".rtf", ".odt", ".ods"
}


def clean_name(text):
    return re.sub(r"\s+", " ", str(text or "")).strip()


def md_table(rows):
    rows = [[clean_name(cell).replace("|", r"\|").replace("\n", "<br>")
             for cell in row] for row in rows]
    rows = [row for row in rows if any(row)]
    if not rows:
        return ""

    width = max(map(len, rows))
    rows = [row + [""] * (width - len(row)) for row in rows]
    output = [
        "| " + " | ".join(rows[0]) + " |",
        "| " + " | ".join(["---"] * width) + " |",
    ]
    output.extend("| " + " | ".join(row) + " |" for row in rows[1:])
    return "\n".join(output)


def convert_pdf(path):
    reader = PdfReader(path)
    parts = []
    for number, page in enumerate(reader.pages, 1):
        parts += [f"## Page {number}", "", (page.extract_text() or "").strip(), ""]
    return "\n".join(parts), len(reader.pages)


def convert_docx(path):
    doc = Document(path)
    parts = []
    for paragraph in doc.paragraphs:
        text = clean_name(paragraph.text)
        if not text:
            continue
        style = paragraph.style.name if paragraph.style else ""
        if style.startswith("Heading"):
            level = re.search(r"\d+", style)
            depth = min(int(level.group()) if level else 2, 6)
            parts.append("#" * depth + " " + text)
        elif "List Bullet" in style:
            parts.append("- " + text)
        elif "List Number" in style:
            parts.append("1. " + text)
        else:
            parts.append(text)
        parts.append("")

    for index, table in enumerate(doc.tables, 1):
        parts += [f"### Table {index}", ""]
        parts.append(md_table([[cell.text for cell in row.cells] for row in table.rows]))
        parts.append("")
    return "\n".join(parts), None


def convert_xlsx(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    parts = []
    for sheet in workbook.worksheets:
        parts += [f"## Sheet: {sheet.title}", ""]
        rows = list(sheet.iter_rows(values_only=True))
        if rows:
            parts.append(md_table([["" if v is None else str(v) for v in row]
                                   for row in rows]))
        parts.append("")
    workbook.close()
    return "\n".join(parts), None


def convert_xls(path):
    workbook = xlrd.open_workbook(path)
    parts = []
    for sheet in workbook.sheets():
        parts += [f"## Sheet: {sheet.name}", ""]
        rows = [[sheet.cell_value(r, c) for c in range(sheet.ncols)]
                for r in range(sheet.nrows)]
        if rows:
            parts.append(md_table(rows))
        parts.append("")
    return "\n".join(parts), None


def convert_pptx(path):
    presentation = Presentation(path)
    parts = []
    for number, slide in enumerate(presentation.slides, 1):
        parts += [f"## Slide {number}", ""]
        for shape in slide.shapes:
            if shape.has_text_frame and clean_name(shape.text):
                parts += [clean_name(shape.text), ""]
            if shape.has_table:
                parts.append(md_table([
                    [cell.text for cell in row.cells]
                    for row in shape.table.rows
                ]))
                parts.append("")
    return "\n".join(parts), len(presentation.slides)


def convert_with_textutil(path):
    result = run(
        ["textutil", "-convert", "txt", "-stdout", str(path)],
        stdout=PIPE, stderr=PIPE, text=True
    )
    if result.returncode != 0:
        raise RuntimeError(result.stderr.strip() or "textutil failed")
    return result.stdout, None


def convert_with_libreoffice(path):
    import tempfile
    with tempfile.TemporaryDirectory() as temp:
        result = run(
            ["libreoffice", "--headless", "--convert-to", "txt:Text",
             "--outdir", temp, str(path)],
            stdout=PIPE, stderr=PIPE, text=True
        )
        converted = Path(temp) / (path.stem + ".txt")
        if result.returncode != 0 or not converted.exists():
            raise RuntimeError(result.stderr.strip() or "LibreOffice conversion failed")
        return converted.read_text(encoding="utf-8", errors="replace"), None


def convert(path):
    ext = path.suffix.lower()
    if ext == ".pdf":
        return convert_pdf(path)
    if ext == ".docx":
        return convert_docx(path)
    if ext == ".xlsx":
        return convert_xlsx(path)
    if ext == ".xls":
        return convert_xls(path)
    if ext == ".pptx":
        return convert_pptx(path)
    if ext in {".doc", ".rtf"}:
        try:
            return convert_with_textutil(path)
        except Exception:
            return convert_with_libreoffice(path)
    if ext in {".ppt", ".odt", ".ods"}:
        return convert_with_libreoffice(path)
    if ext == ".txt":
        return path.read_text(encoding="utf-8", errors="replace"), None
    raise ValueError(f"Unsupported file type: {ext}")


def main():
    if not SOURCE.exists():
        print(f"Source folder not found: {SOURCE}")
        return

    DEST.mkdir(parents=True, exist_ok=True)
    files = sorted(
        p for p in SOURCE.rglob("*")
        if p.is_file()
        and p.suffix.lower() in SUPPORTED
        and DEST not in p.parents
    )

    print(f"Found {len(files)} files to convert.")
    converted = 0
    failed = 0

    for path in files:
        output = DEST / (path.stem + ".md")
        if output.exists():
            print(f"SKIP  {path.name} — Markdown file already exists")
            continue

        try:
            text, count = convert(path)
            header = [
                f"# {path.stem.replace('-', ' ')}",
                "",
                f"- Original file: `{path.name}`",
                "- Source: Programme on a Plate, Scouts South Africa Wiki",
            ]
            if count is not None:
                header.append(f"- Pages/slides: {count}")
            header += ["", "---", ""]
            output.write_text(
                "\n".join(header) + text.strip() + "\n",
                encoding="utf-8"
            )
            print(f"OK    {path.name} -> {output.name}")
            converted += 1
        except Exception as error:
            print(f"FAIL  {path.name}: {error}")
            failed += 1

    print(f"\nFinished: {converted} converted, {failed} failed.")
    print(f"Markdown files are in: {DEST}")


if __name__ == "__main__":
    main()
